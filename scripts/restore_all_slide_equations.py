#!/usr/bin/env python3
"""Replace ALL 22 embedded mathematical graphics in the 44-slide deck with native OfficeMath.

Bounded, source-to-formula-mapped conversion:
- Convert only known original mathematical image SHA-256 prefixes, never charts/logos.
- Every expression is transcribed from the corresponding image in the same Git commit.
- Preserve original caption/text, original equation slot, slide order and non-math images.
- Patch only changed slide*.xml in source ZIP; untouched PowerPoint parts byte-identical.
- Retain original Git history and emit a provenance manifest with formula transcriptions.

Reconstructed OMML is editable but is not falsely described as the unavailable original
author-entered equation object. Does not alter experimental calculations or metadata.
"""
from __future__ import annotations
import sys,json,zipfile,hashlib
from copy import deepcopy
from pathlib import Path
from collections import Counter
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE as Type
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from pptx.oxml.ns import qn

M="http://schemas.openxmlformats.org/officeDocument/2006/math"
A14="http://schemas.microsoft.com/office/drawing/2010/main"
DRAW="http://schemas.openxmlformats.org/drawingml/2006/main"
R="{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
# DSL: text; ("sub",base,sub); ("sup",base,sup);
# ("ss",base,sub,sup); ("f",numerator,denominator);
# ("rad",radicand); ("n",operator,lower,upper,body);
# ("acc",base,"bar"|"hat").
def S(base,sub):return ("sub",base,sub)
def U(base,sup):return ("sup",base,sup)
def B(base,sub,sup):return ("ss",base,sub,sup)
def F(num,den):return ("f",num,den)
def RAD(inside):return ("rad",inside)
def N(op,lo,up,body):return ("n",op,lo,up,body)
def ACC(base,accent):return ("acc",base,accent)
def C(*parts):return list(parts)

# Exact reading of the original IMAGE, not invented numerical results.
# The dictionary keys are SHA256 prefixes of the corresponding original media.
MATH_BY_SLIDE={
4:[
 ("4c0726008","MSE over d dimensions",C(S("L","t")," = MSE(",S("X","t"),", ",ACC(S("X","t"),"hat"),") = ",F("1","d"),N("∑","j=1","d",C("(",S("X","t,j")," − ",ACC(S("X","t,j"),"hat"),U(")","2")))),"white")
],
9:[
 ("bebd58d6b","Kronecker delta",C(S("δ","ij")," = 1 (i = j),     ",S("δ","ij")," = 0 (i ≠ j)"),"dark")
],
10:[
 ("1b76839ae","Covariance and Kronecker identity",C("Cov(",S("Z","i"),",",S("Z","j"),") = E[",S("Z","i"),S("Z","j"),"] = ",S("δ","ij")),"dark")
],
11:[
 ("91eef13af","Partial derivative identity",C(F("∂"+ "wᵢ","∂"+"wⱼ")," = ",S("δ","ij")),"dark"),
 ("6ef44f2ec","Variance identity",C("Var(",S("Z","i"),") = E[(",S("Z","i")," − E[",S("Z","i"),"]",U(")","2"),"]"),"dark"),
 ("689cb93a6","Online update",C(S("w","i"),"(t+1) = ",S("w","i"),"(t) + ηe(t)",S("x","i"),"(t)"),"dark")
],
12:[
 ("a26836a86","Vanishing wavelet moments",C("∫ ",U("t","m"),"ψ(t)dt = 0,     m = 0,…,N − 1"),"dark")
],
13:[
 ("90a2256e1","Linear transformation and covariance",C("Y = AZ,    Cov(Z) = I    ⇒    Cov(Y) = A",U("A","T")),"dark"),
 ("918bcb4e3","Correlation construction",C(S("Y","1")," = ",S("Z","1"),",    ",S("Y","2")," = a",S("Z","1")," + ",S("Z","2"),",    Cov(",S("Y","1"),",",S("Y","2"),") = a"),"dark")
],
20:[
 ("890c6953a","CSP spatial filter",C("Z = WE"),"white"),
 ("c72320940","Normalised CSP log-variance feature",C(S("f","j")," = log", "(",F(C("Var(",S("Z","j"),")"), N("∑","k=1","2m",C("Var(",S("Z","k"),")"))),")"),"white"),
 ("54ff4741a","PCA projection and first component",C(S("s","i")," = ",U("P","T"),"(",S("f","i")," − ",S("μ","cal"),"),     ",S("X","i")," = ",S("s","i,1")),"white")
],
21:[
 ("5d51b3a3e","Representation and adaptation",C(S("f","t")," = ",S("ϕ","0"),"(",S("E","t"),"),    ",S("q","t")," = ",U(S("P","0"),"T"),"(",S("f","t")," − ",S("u","0"),"),    ",S("ŷ","t")," = ",S("g",S("θ","t")),"(",S("f","t"),")"),"dark")
],
22:[
 ("18da8bf0e","EWMA statistic and innovation",C(S("z","i")," = λ",S("x","i")," + (1−λ)",S("z","i−1"),",    ",S("e","i")," = ",S("x","i")," − ",S("z","i−1")),"white"),
 ("ee619e437","EWMA upper and lower control limits",C(S("UCL","i")," = ",S("z","i−1")," + L",S("σ","e,i−1"),",    ",S("LCL","i")," = ",S("z","i−1")," − L",S("σ","e,i−1")),"white"),
 ("e345c7f35","Variance EWMA recurrence",C(B("σ","e,i","2")," = ",S("α","v"),U(S("e","i"),"2")," + (1−",S("α","v"),")",B("σ","e,i−1","2")),"white"),
 ("df90e7159","Two-sample Hotelling T squared",C(U("T","2")," = ",F(C(S("n","1"),S("n","2")),C(S("n","1"),"+",S("n","2"))),"(",ACC(S("x","1"),"bar")," − ",ACC(S("x","2"),"bar"),U(")","T"),U(S("S","p"),"−1"),"(",ACC(S("x","1"),"bar")," − ",ACC(S("x","2"),"bar"),")"),"white"),
 ("97f0972b5","Detection threshold and degrees of freedom",C("p < ",S("α","test"),",     2H − d − 1 > 0"),"white")
],
23:[
 ("d162d37a3","One-sample online T squared",C(B("T","t","2")," = (",S("z","t")," − ",S("μ","0"),U(")","T"),U(S("Σ","z,t"),"−1"),"(",S("z","t")," − ",S("μ","0"),")"),"dark"),
 ("f0df7120c","Pooled two-sample Hotelling",C(U("T","2")," = ",F(C(S("n","1"),S("n","2")),C(S("n","1")," + ",S("n","2"))),"(",ACC(S("x","1"),"bar")," − ",ACC(S("x","2"),"bar"),U(")","T"),U(S("S","p"),"−1"),"(",ACC(S("x","1"),"bar")," − ",ACC(S("x","2"),"bar"),")"),"dark")
],
26:[
 ("938853d14","Unpooled covariance weighted mean difference",C("Q = (",ACC(S("x","1"),"bar")," − ",ACC(S("x","2"),"bar"),U(")","T"),U(C("(",F(S("S","1"),S("n","1"))," + ",F(S("S","2"),S("n","2")),")"),"−1"),"(",ACC(S("x","1"),"bar")," − ",ACC(S("x","2"),"bar"),")"),"dark")
],
36:[
 ("2e5addd00","CUSUM auxiliary quantity",C(S("Q","t")," = d",S("L","t")),"cyan")
]
}
EXPECTED={i:len(v) for i,v in MATH_BY_SLIDE.items()}
assert sum(EXPECTED.values())==22,EXPECTED

def n(kind, parent):
    return etree.SubElement(parent,f"{{{M}}}{kind}")
def atom_text(node,txt):
    r=n("r",node)
    # Some PowerPoint-compatible mathematical run-property combinations
    # disappear in LibreOffice. Do not emit redundant m:rPr for ordinary runs.
    n("t",r).text=str(txt)
def render(parent,obj):
    if isinstance(obj,str):
        atom_text(parent,obj);return
    if isinstance(obj,list):
        for e in obj:render(parent,e)
        return
    if not isinstance(obj,tuple):raise TypeError(obj)
    typ=obj[0]
    if typ in ("sub","sup","ss"):
        e=n({"sub":"sSub","sup":"sSup","ss":"sSubSup"}[typ],parent)
        base=n("e",e);render(base,obj[1])
        if typ in ("sub","ss"):
            sub=n("sub",e);render(sub,obj[2])
        if typ in ("sup","ss"):
            sup=n("sup",e);render(sup,obj[2 if typ=="sup" else 3])
    elif typ=="f":
        e=n("f",parent)
        for tag,value in (("num",obj[1]),("den",obj[2])):
            q=n(tag,e);render(q,value)
    elif typ=="rad":
        e=n("rad",parent)
        p=n("radPr",e);dh=n("degHide",p);dh.set(f"{{{M}}}val","1")
        n("deg",e)
        q=n("e",e);render(q,obj[1])
    elif typ=="n":
        e=n("nary",parent)
        pr=n("naryPr",e);ch=n("chr",pr);ch.set(f"{{{M}}}val",obj[1])
        for tag,value in (("sub",obj[2]),("sup",obj[3]),("e",obj[4])):
            q=n(tag,e);render(q,value)
    elif typ=="acc":
        e=n("acc",parent)
        pr=n("accPr",e);ch=n("chr",pr)
        ch.set(f"{{{M}}}val",{"bar":"̄","hat":"̂"}[obj[2]])
        q=n("e",e);render(q,obj[1])
    else:raise ValueError(obj)

def new_math_shape(slide,name,rect,expr,color,font_size):
    # Existing freeform graphics are grouped. Create the mathematical equivalent
    # at the group bounding rectangle and replace that group at the same z-order.
    box=slide.shapes.add_textbox(*rect)
    box.name="EDITABLE OfficeMath | "+name
    tf=box.text_frame
    tf.clear();tf.word_wrap=False
    tf.margin_left=tf.margin_right=0
    tf.margin_top=tf.margin_bottom=0
    p=tf.paragraphs[0]
    p.alignment=PP_ALIGN.CENTER
    end=p._p.get_or_add_endParaRPr();end.set("sz",str(font_size*100))
    end.set("lang","en-GB")
    fg=etree.SubElement(end,f"{{{DRAW}}}solidFill")
    rgb=etree.SubElement(fg,f"{{{DRAW}}}srgbClr")
    rgb.set("val",{"white":"14233B","dark":"14233B","cyan":"075F57"}[color])
    if color in ("white","cyan"):
        # LibreOffice's imported OfficeMath uses black by default even when the
        # paragraph requests white. Small contrasting equation backing keeps
        # equations readable in PowerPoint and LibreOffice while preserving
        # their exact location, editability and the surrounding slide design.
        sppr=box._element.xpath('./*[local-name()="spPr"]')[0]
        for el in list(sppr):
            if el.tag in (f"{{{DRAW}}}noFill",f"{{{DRAW}}}solidFill"):
                sppr.remove(el)
        bg=etree.SubElement(sppr,f"{{{DRAW}}}solidFill")
        color_ref=etree.SubElement(bg,f"{{{DRAW}}}srgbClr")
        color_ref.set("val","DDEBF3" if color=="white" else "D5F6F1")
    mat=etree.Element(f"{{{A14}}}m",nsmap={"a14":A14,"m":M})
    para=n("oMathPara",mat)
    eq=n("oMath",para)
    render(eq,expr)
    end.addprevious(mat)
    return box

def text(sl):
    return [sh.text for sh in sl.shapes if sh.has_text_frame and not sh.name.startswith(("EDITABLE OfficeMath |", "EDITABLE MATH |"))]
def findhash(slide,shape):
    if shape.shape_type!=Type.GROUP or len(shape.shapes)!=1:return None
    refs=shape.shapes[0]._element.xpath('.//*[local-name()="blip"]')
    if len(refs)!=1:return None
    rid=refs[0].get(R)
    if not rid:return None
    return hashlib.sha256(slide.part.related_part(rid).blob).hexdigest()

def split_prior_math(slide,number:int):
    """Retain each existing native OMML expression but place it in its own editable box.

    Multiple OfficeMath paragraphs inside one PowerPoint textbox can render as
    nothing in alternative viewers; this uses unchanged original m:oMath subtrees.
    """
    restored=0
    for shape in list(slide.shapes):
        if not shape.name.startswith("EDITABLE MATH |"):continue
        paras=shape._element.xpath('./*[local-name()="txBody"]/*[local-name()="p"]')
        if len(paras)<2:continue
        source=[etree.tostring(el) for el in shape._element.xpath('.//*[local-name()="oMath"]')]
        parent=shape._element.getparent()
        index=parent.index(shape._element)
        nrows=len(paras)
        # Slide 7 has a caption directly above the equation and a notation key below;
        # shift rows inside that panel so neither text block is covered.
        start=int(Inches(6.82)) if number==7 else int(shape.top)
        bottom=int(Inches(8.18)) if number==7 else int(shape.top+shape.height-Inches(.30))
        height=(bottom-start)//nrows
        for j in range(nrows):
            cloned=deepcopy(shape._element)
            textbody=cloned.xpath('./*[local-name()="txBody"]')[0]
            rows=cloned.xpath('./*[local-name()="txBody"]/*[local-name()="p"]')
            for k,row in enumerate(rows):
                if k!=j:textbody.remove(row)
            xfrm=cloned.xpath('./*[local-name()="spPr"]/*[local-name()="xfrm"]')[0]
            off=xfrm.xpath('./*[local-name()="off"]')[0]
            ext=xfrm.xpath('./*[local-name()="ext"]')[0]
            off.set("y",str(start+j*height))
            ext.set("cy",str(height))
            c=cloned.xpath('./*[local-name()="nvSpPr"]/*[local-name()="cNvPr"]')[0]
            c.set("id",str(6000+number*10+j))
            c.set("name",shape.name+" | ROW "+str(j+1))
            # Lower font size on dense four-row slide 7, preserving the original
            # equation symbol structures and their provenance.
            if number==7:
                for end in cloned.xpath('.//*[local-name()="endParaRPr"]'):
                    end.set("sz","1300")
            parent.insert(index+j,cloned)
        parent.remove(shape._element)
        after=[etree.tostring(el) for el in slide._element.xpath('.//*[local-name()="oMath"]')]
        for item in source:
            assert item in after,"A pre-existing Park equation object was altered"
        restored+=nrows
    return restored

def build(src:Path,dst:Path,report:Path):
    if src.resolve()==dst.resolve():raise ValueError("Never overwrite the source")
    prs=Presentation(str(src))
    assert len(prs.slides)==44
    before=[text(sl) for sl in prs.slides]
    old_native={
        n:[etree.tostring(m) for m in prs.slides[n-1]._element.xpath('.//*[local-name()="oMath"]')]
        for n in (7,8)
    }
    prior_split=sum(split_prior_math(prs.slides[n-1],n) for n in (7,8))
    assert prior_split==10, f"Expected ten pre-existing editable Park equations, found {prior_split}"
    for n in (7,8):
        assert len(prs.slides[n-1]._element.xpath('.//*[local-name()="oMath"]'))==len(old_native[n])
    slide_xml={
        f"ppt/slides/slide{n}.xml":etree.tostring(prs.slides[n-1]._element,encoding="UTF-8",xml_declaration=True,standalone=True)
        for n in (7,8)
    }
    original_count={i:len(sl._element.xpath('.//*[local-name()="oMath"]')) for i,sl in enumerate(prs.slides,1)}
    assert original_count[7]>=4 and original_count[8]>=6,"Main should already include the prior native equation restoration"
    allchanges=[]
    for number,items in MATH_BY_SLIDE.items():
        slide=prs.slides[number-1]
        original_omml=original_count[number]
        found=[]
        for group in list(slide.shapes):
            digest=findhash(slide,group)
            if not digest:continue
            matches=[i for i,spec in enumerate(items) if digest.startswith(spec[0])]
            if not matches:continue
            if len(matches)!=1:raise AssertionError((number,digest,matches))
            idx=matches[0]
            if idx in found:raise AssertionError((number,idx,"duplicate media"))
            sha,name,expr,color=items[idx]
            rect=(group.left,group.top,group.width,group.height)
            # Derived from original equation group height. Keep original placement.
            inch_h=group.height/914400.0
            size=max(18,min(31,int(inch_h*27)))
            if number in (13,21):size=min(size,23)
            if number in (4,20,22):size=min(size,26)
            box=new_math_shape(slide,name,rect,expr,color,size)
            parent=group._element.getparent()
            parent.replace(group._element,box._element)
            found.append(idx)
            allchanges.append({
                "slide":number,"source_image_sha256":digest,"formula_name":name,
                "bounding_box_emu":list(map(int,rect)),
                "math_object":"native PowerPoint OfficeMath (m:oMath)",
                "equation_object_sha256":hashlib.sha256(etree.tostring(box._element)).hexdigest(),
                "font_color":color,
                "new_native_math_count":len(box._element.xpath('.//*[local-name()="oMath"]')),
                "representation":"human-transcribed exact displayed formula from committed image; editable OMML, not original binary equation object"
            })
        if sorted(found)!=list(range(len(items))):
            raise AssertionError(f"Slide {number}: expected {len(items)} original images, replaced indices {found}")
        assert len(slide._element.xpath('.//*[local-name()="oMath"]'))==original_omml+len(items)
        assert text(slide)==before[number-1],f"Slide {number}: original text changed"
        slide_xml[f"ppt/slides/slide{number}.xml"]=etree.tostring(slide._element,encoding="UTF-8",xml_declaration=True,standalone=True)
        print("CONVERTED",number,len(items),"editable formulas")
    assert len(allchanges)==22
    assert all(text(sl)==old for sl,old in zip(prs.slides,before))
    dst.parent.mkdir(parents=True,exist_ok=True)
    # Zip surgery preserves all untouched original PPTX ZIP members exactly.
    with zipfile.ZipFile(src,"r") as zin,zipfile.ZipFile(dst,"w") as zout:
        assert set(slide_xml).issubset(zin.namelist())
        for entry in zin.infolist():
            zout.writestr(entry,slide_xml.get(entry.filename,zin.read(entry.filename)))
    with zipfile.ZipFile(src) as zin,zipfile.ZipFile(dst) as zout:
        assert zin.namelist()==zout.namelist()
        for name in zin.namelist():
            if name not in slide_xml:
                assert zin.read(name)==zout.read(name),f"NON-MATH ZIP MEMBER MODIFIED: {name}"
        for slide in range(1,45):
            member=f"ppt/slides/slide{slide}.xml"
            if member not in slide_xml:
                assert zin.read(member)==zout.read(member),f"UNEXPECTED CHANGED SLIDE {slide}"
    check=Presentation(str(dst))
    counts={i:len(check.slides[i-1]._element.xpath('.//*[local-name()="oMath"]')) for i in range(1,45)}
    for n_,k in EXPECTED.items():assert counts[n_]==original_count[n_]+k
    for i in range(1,45):
        if i not in EXPECTED:assert counts[i]==original_count[i]
    assert sum(counts.values())==sum(original_count.values())+22
    assert all(text(sl)==old for sl,old in zip(check.slides,before))
    report.parent.mkdir(parents=True,exist_ok=True)
    report.write_text(json.dumps({
      "input_sha256":hashlib.sha256(src.read_bytes()).hexdigest(),
      "output_sha256":hashlib.sha256(dst.read_bytes()).hexdigest(),
      "source_slide_count":len(prs.slides),
      "new_native_math_objects":22,
      "prior_math_objects_preserved":sum(original_count.values()),
      "total_native_math_objects":sum(counts.values()),
      "untouched_zip_parts_byte_identical":True,
      "changed_slide_numbers":sorted(set(EXPECTED)|{7,8}),
      "conversion_provenance":allchanges,
    },ensure_ascii=False,indent=2),encoding="utf-8")
    print("PASS 44 slides; 22 replacement native equations; 10 existing native equations retained and split for visibility.")
    print("PASS every ZIP entry outside the targeted slide XMLs is byte-for-byte identical.")
    print("PASS logos, plots, timeline graphics, experiments and original text preserved.")
    print("PROVENANCE",report)
if __name__=="__main__":
    src,dst,report=map(Path,sys.argv[1:4])
    build(src,dst,report)
