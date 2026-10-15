#!/usr/bin/env python3
"""Make all slide formulas legible as editable mathematical text; retain native OMML sources."""
import sys,json,zipfile,hashlib
from pathlib import Path
from pptx import Presentation
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from lxml import etree

source=Path(sys.argv[1]);target=Path(sys.argv[2]);manifest=Path(sys.argv[3])
rows=json.loads(manifest.read_text(encoding="utf-8"))["rows"]
deck=Presentation(str(source))
assert len(deck.slides)==44
assert sum(len(s._element.xpath('.//*[local-name()="oMath"]')) for s in deck.slides)==32
base_text=[
    [sh.text for sh in slide.shapes if sh.has_text_frame and
     not sh.name.startswith(("EDITABLE MATH |","EDITABLE OfficeMath |"))]
    for slide in deck.slides
]
changed={}
found=[]
for slide_num in range(1,45):
    slide=deck.slides[slide_num-1]
    requested=[r for r in rows if r[0]==slide_num]
    for sh in list(slide.shapes):
        if not sh.name.startswith(("EDITABLE MATH |","EDITABLE OfficeMath |")):
            continue
        label=sh.name.split(" | ",1)[1]
        matches=[r for r in requested if r[1]==label]
        assert len(matches)==1,(slide_num,label)
        r=matches[0]
        _,_,eq,pt,x,y,w,h=r
        sh.left=Inches(22.5)
        sh.name="OMML SOURCE (retained; via Selection Pane) | "+label
        box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
        box.name="COMPACT EDITABLE FORMULA | "+label
        tf=box.text_frame;tf.clear();tf.word_wrap=False
        tf.margin_left=tf.margin_right=Inches(.04)
        tf.margin_top=tf.margin_bottom=0
        tf.vertical_anchor=MSO_ANCHOR.MIDDLE
        paragraph=tf.paragraphs[0]
        paragraph.alignment=PP_ALIGN.CENTER if slide_num in (7,8,9,10,11,12,13,23,26) else PP_ALIGN.LEFT
        run=paragraph.add_run()
        run.text=eq;run.font.name="Cambria Math";run.font.size=Pt(pt)
        run.font.color.rgb=RGBColor(226,237,246) if slide_num in (4,20,22) else RGBColor(19,36,60)
        if slide_num==36:run.font.color.rgb=RGBColor(10,121,122)
        found.append((slide_num,label))
    if requested:
        changed[f"ppt/slides/slide{slide_num}.xml"]=etree.tostring(
            slide._element,encoding="UTF-8",xml_declaration=True,standalone=True)
assert len(found)==len(rows)==32,(len(found),len(rows))
for i,slide in enumerate(deck.slides):
    now=[
        sh.text for sh in slide.shapes if sh.has_text_frame and
        not sh.name.startswith(("OMML SOURCE (retained; via Selection Pane) |","COMPACT EDITABLE FORMULA |"))]
    assert now==base_text[i],("existing slide text changed",i+1)
target.parent.mkdir(parents=True,exist_ok=True)
with zipfile.ZipFile(source,"r") as zi,zipfile.ZipFile(target,"w") as zo:
    for part in zi.infolist():
        zo.writestr(part,changed.get(part.filename,zi.read(part.filename)))
with zipfile.ZipFile(source) as zi,zipfile.ZipFile(target) as zo:
    assert zi.namelist()==zo.namelist()
    for part in zi.namelist():
        if part not in changed:assert zi.read(part)==zo.read(part),part
test=Presentation(str(target))
native=sum(len(s._element.xpath('.//*[local-name()="oMath"]')) for s in test.slides)
visible=sum(sh.name.startswith("COMPACT EDITABLE FORMULA |") for s in test.slides for sh in s.shapes)
assert len(test.slides)==44 and native==visible==32,(len(test.slides),native,visible)
print("PASS: 44 slides, 32 visible editable mathematical text objects, 32 original native OfficeMath objects retained")
print("PASS: non-formula text and all other PPTX ZIP components unchanged")
print("MODIFIED",sorted({r[0] for r in rows}))
print("SHA256",hashlib.sha256(target.read_bytes()).hexdigest())
