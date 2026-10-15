#!/usr/bin/env python3
"""Make an alternate PowerPoint with standard picture objects for existing equations.

The original deck is never overwritten. All equation pixels come from the
existing PPTX ZIP; no formula transcription or scientific recomputation occurs.
"""
from __future__ import annotations

import argparse
from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
DEFAULT_SLIDES = {9, 10, 11, 12, 13, 20, 22, 23, 26}


def convert(original: Path, target: Path, slide_numbers: set[int]) -> list[dict]:
    if original.resolve() == target.resolve():
        raise ValueError("Refusing to overwrite the original research presentation")
    presentation = Presentation(str(original))
    changes = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        if slide_number not in slide_numbers:
            continue
        for shape in list(slide.shapes):
            if shape.shape_type != MSO_SHAPE_TYPE.GROUP:
                continue
            if len(shape.shapes) != 1:
                continue
            child = shape.shapes[0]
            if child.shape_type != MSO_SHAPE_TYPE.FREEFORM:
                continue
            blips = child._element.xpath(".//*[local-name()='blip']")
            if len(blips) != 1:
                continue
            if child._element.xpath(".//*[local-name()='srcRect']"):
                raise ValueError(f"Slide {slide_number}: cropped image needs special handling")
            reference = blips[0].get(R_EMBED)
            relation = slide.part.rels.get(reference) if reference else None
            if relation is None or relation.is_external:
                raise ValueError(f"Slide {slide_number}: missing embedded equation image")
            image_bytes = relation.target_part.blob
            if relation.target_part.content_type != "image/png":
                continue
            # Retain exact embedded raster bytes and preserve drawing order.
            picture = slide.shapes.add_picture(
                BytesIO(image_bytes), shape.left, shape.top, shape.width, shape.height
            )
            parent = shape._element.getparent()
            parent.replace(shape._element, picture._element)
            changes.append({
                "slide": slide_number,
                "image_bytes": len(image_bytes),
                "position_emu": [picture.left, picture.top, picture.width, picture.height],
            })
    if not changes:
        raise ValueError("No image-filled equation groups found; no output created")
    target.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(target))
    return changes


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("original", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument(
        "--slides",
        default=",".join(map(str, sorted(DEFAULT_SLIDES))),
        help="Comma-separated physical slide numbers",
    )
    args = parser.parse_args()
    slides = {int(value.strip()) for value in args.slides.split(",") if value.strip()}
    changes = convert(args.original, args.output, slides)
    for item in changes:
        print("REWRAPPED", item)
    print(f"Rewrapped {len(changes)} existing equation images without recreating formulas")
    print(f"Compatibility copy: {args.output}")


if __name__ == "__main__":
    main()
