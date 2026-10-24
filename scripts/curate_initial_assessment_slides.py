"""Produce a concise assessment copy, preserving all original source files.

Only remove/reorder slides in the DERIVED presentation. Never write the source.
The new WP1 readiness slide is marked as an assessment synthesis, not an
experiment. Original equations, illustrations and data are not rewritten.
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "presentations/EWMA_Covariate_Shift_Progress_Review_Haocheng_Lin.pptx"
OUTPUT = ROOT / "presentations/EWMA_Covariate_Shift_Initial_Assessment_Curated.pptx"
# One-based original slide indices. Preserve original detailed deck separately.
KEEP = [
    1, 2, 3, 5, 8, 9, 25, 11, 32, 33, 39, 36, 13, 43, 46, 47, 48, 12, 52, 54
]
DESTINATION_AFTER = 18   # Insert WP1 assessment summary after the evidence ledger.

def textbox(slide,left,top,width,height,text,pt=14,bold=False,color=(28,45,68),align=None):
    sh=slide.shapes.add_textbox(left,top,width,height)
    tf=sh.text_frame;tf.word_wrap=True
    tf.margin_top=Pt(2);tf.margin_left=Pt(3);tf.margin_bottom=Pt(1)
    p=tf.paragraphs[0];p.text=text;p.font.name="Aptos";p.font.size=Pt(pt);p.font.bold=bold
    p.font.color.rgb=RGBColor(*color)
    if align is not None:p.alignment=align
    return sh

def rectangle(slide,left,top,width,height,color):
    sh=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,left,top,width,height)
    sh.fill.solid();sh.fill.fore_color.rgb=RGBColor(*color)
    sh.line.fill.background();return sh

def add_wp1_slide(prs):
    s=prs.slides.add_slide(prs.slide_layouts[6])
    w,h=prs.slide_width,prs.slide_height
    rectangle(s,0,0,w,h, (249,251,253))
    rectangle(s,0,0,w, Inches(.17),(20,58,83))
    textbox(s,Inches(.58),Inches(.38),w-Inches(1),Inches(.48),
            "WP1 — INITIAL ASSESSMENT CHECKPOINT",25,True,(19,55,81))
    textbox(s,Inches(.58),Inches(.90),w-Inches(1),Inches(.47),
            "24 October 2026 | Theoretical Formulation & Architecture Definition (Months 1–3)",
            12,False,(75,95,112))
    rows=[
      ("FORMULATION", "Substantial foundation", "Covariate vs concept drift; EWMA, Hotelling, Park and ABCD maths."),
      ("STABILITY–PLASTICITY", "Policy laboratory", "Six rules in code; long-term retention and energy trade-offs not yet validated."),
      ("GUARANTEES / BUDGETS", "Under investigation", "False-alarm multiplicity identified; calibrated ARL and resource bounds open."),
      ("CPD–CL–TWIN ARCHITECTURE", "Modular prototype", "Detection-to-policy interfaces exist; CL/twin validation contract remains proposed."),
      ("ROBUSTNESS", "Evidence-rich audit", "B01 count conflict, invalid Hotelling settings and CPSS calibration limits documented.")
    ]
    top=Inches(1.63);rowh=(h-Inches(2.47))/5
    for i,(name,status,desc) in enumerate(rows):
        y=top+i*rowh
        if i%2==0:rectangle(s,Inches(.50),y,w-Inches(1),rowh,(236,242,247))
        rectangle(s,Inches(.54),y+Inches(.16),Inches(.055),rowh-Inches(.32),(20,119,126))
        textbox(s,Inches(.70),y+Inches(.08), Inches(2.55),rowh-Inches(.10),name,13,True)
        textbox(s,Inches(3.15),y+Inches(.08), Inches(2.35),rowh-Inches(.10),status,12,True,(20,110,120))
        textbox(s,Inches(5.45),y+Inches(.07),w-Inches(6.1),rowh-Inches(.12),desc,12)
    textbox(s,Inches(.58),h-Inches(.54),w-Inches(1.15),Inches(.29),
            "Assessment synthesis; supporting evidence in main.tex, Notion WP1 and GitHub outputs. Not a validated end-to-end system.",
            10,False,(82,100,115))
    return s

def main():
    prs=Presentation(SOURCE)
    if len(prs.slides)!=55:
        raise RuntimeError(f"Deck changed: expected 55 original slides, got {len(prs.slides)}")
    original=list(prs.slides._sldIdLst)
    for i in reversed(range(len(original))):
        if i+1 not in KEEP:
            node=original[i]
            prs.part.drop_rel(node.rId)
            prs.slides._sldIdLst.remove(node)
    retained=list(prs.slides._sldIdLst)
    for node in retained:prs.slides._sldIdLst.remove(node)
    by_source={source:node for source,node in zip(sorted(KEEP),retained)}
    for source in KEEP:prs.slides._sldIdLst.append(by_source[source])
    # The cover is a dated snapshot; update only the date line on this derived copy.
    for shape in prs.slides[0].shapes:
        if shape.has_text_frame and "20th October 2026" in shape.text:
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if "20th October 2026" in run.text:
                        run.text=run.text.replace("20th October 2026","24 October 2026")
    summary=add_wp1_slide(prs)
    summary_node=prs.slides._sldIdLst[-1]
    prs.slides._sldIdLst.remove(summary_node)
    prs.slides._sldIdLst.insert(DESTINATION_AFTER,summary_node)
    prs.save(OUTPUT)
    test=Presentation(OUTPUT)
    assert len(test.slides)==len(KEEP)+1, (len(test.slides),len(KEEP))
    print(f"Curated deck {len(test.slides)} slides from 55; original unchanged")
    print(f"Output: {OUTPUT}")

if __name__=="__main__":
    main()
