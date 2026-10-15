#!/usr/bin/env python3
"""Native, editable OfficeMath repair for Park slides 7–8 (not image rendering)."""
import hashlib,sys,zipfile
from pathlib import Path
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE as Type
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches,Pt

M="http://schemas.openxmlformats.org/officeDocument/2006/math"
A14="http://schemas.microsoft.com/office/drawing/2010/main"
R="{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"

def elem(kind,parent):
    x=etree.SubElement(parent,f"{{{M}}}{kind}")
    return x
def run(parent,value):
    r=elem("r",parent)
    elem("t",r).text=value
def pieces(parent,parts):
    for item in parts:
        if isinstance(item,str):run(parent,item)
        elif item[0]=="frac":
            f=elem("f",parent)
            n=elem("num",f); pieces(n,item[1])
            d=elem("den",f); pieces(d,item[2])
        elif item[0]=="ss":
            s=elem("sSubSup",parent)
            for tag,value in zip(("e","sub","sup"),item[1:]):
                a=elem(tag,s);run(a,value)
        else:raise ValueError(item)
def add_math(slide,name,rect,lines,font_size):
    box=slide.shapes.add_textbox(*rect)
    box.name="EDITABLE MATH | "+name
    tf=box.text_frame;tf.clear();tf.word_wrap=False
    tf.margin_left=tf.margin_right=Inches(.05)
    tf.margin_top=tf.margin_bottom=Inches(.04)
    for i,components in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment=PP_ALIGN.CENTER
        p.space_after=Pt(8 if len(lines)<=2 else 4)
        end=p._p.get_or_add_endParaRPr();end.set("sz",str(font_size*100))
        wrapper=etree.Element(f"{{{A14}}}m",nsmap={"a14":A14,"m":M})
        omath=elem("oMath",elem("oMathPara",wrapper))
        pieces(omath,components)
        end.addprevious(wrapper)
    return box
def text(slide):
    return "\n".join(s.text for s in slide.shapes if s.has_text_frame and not s.name.startswith("EDITABLE MATH | "))
def media(slide):
    out=[]
    for sh in slide.shapes:
        try: typ=sh.shape_type
        except NotImplementedError:typ=None
        children=list(sh.shapes) if typ==Type.GROUP else [sh]
        for c in children:
            for blip in c._element.xpath('.//*[local-name()="blip"]'):
                rid=blip.get(R)
                if rid:
                    out.append(hashlib.sha256(slide.part.related_part(rid).blob).hexdigest())
    return out
def repair(src,dst):
    if src.resolve()==dst.resolve():raise ValueError("Refuse to overwrite input")
    deck=Presentation(str(src))
    assert len(deck.slides)==44
    original_text=[text(s) for s in deck.slides]
    other_media=[media(s) for s in deck.slides]
    s7,s8=deck.slides[6],deck.slides[7]
    assert "CODE-TO-MANUSCRIPT BOUNDARY" in original_text[6]
    assert "From Channelwise Mean-CUSUM" in original_text[7]
    assert "STEP 1: Aggregation Across Channels" in original_text[7]
    assert not s7._element.xpath('.//*[local-name()="oMath"]')
    assert not s8._element.xpath('.//*[local-name()="oMath"]')
    # This specific JPEG is the unrelated landscape image occupying the equation slot.
    # Never remove an unrecognised asset on a newly edited presentation.
    scenery="c80c612b0b9363d50925ec384064c550f6597b3c96b27f165517b02165cc061e"
    found=[]
    for sh in s7.shapes:
        if sh.shape_type!=Type.GROUP:continue
        imgs=[]
        for c in sh.shapes:
            for blip in c._element.xpath('.//*[local-name()="blip"]'):
                rid=blip.get(R)
                if rid:imgs.append(hashlib.sha256(s7.part.related_part(rid).blob).hexdigest())
        if scenery in imgs:found.append(sh)
    assert len(found)==1,"Slide 7 source asset changed: abort without altering anything"
    bad=found[0];rect=(bad.left,bad.top,bad.width,bad.height)
    bad._element.getparent().remove(bad._element)
    # Equivalent decomposition of Park et al. (2023) Eq. (2.2):
    # nu = |weighted left sum - weighted right sum| / channel mean.
    # m:sSubSup and m:f are actual OfficeMath structures.
    add_math(s7,"Park 2023 Eq 2.2 normalized CUSUM",rect,[
        [[ "ss","ν","t","(d)"]," = ",["frac",["|Aₜ⁽ᵈ⁾ − Bₜ⁽ᵈ⁾|"],["X̄⁽ᵈ⁾"]]],
        ["Aₜ⁽ᵈ⁾ = √((T−t)/(Tt)) ∑ᵤ₌₀ᵗ⁻¹ ",["ss","X","u,T","(d)"]],
        ["Bₜ⁽ᵈ⁾ = √(t/[T(T−t)]) ∑ᵤ₌ₜᵀ⁻¹ ",["ss","X","u,T","(d)"]],
        ["X̄⁽ᵈ⁾ = (1/T) ∑ᵤ₌₀ᵀ⁻¹ ",["ss","X","u,T","(d)"]],
    ],16)
    panels=[s for s in s8.shapes if s.shape_type==Type.GROUP]
    assert len(panels)==2 and all(p.width>Inches(7) for p in panels)
    def rect_for(g):return (g.left+Inches(.36),g.top+Inches(1.55),g.width-Inches(.72),Inches(3.50))
    add_math(s8,"Park 2023 pointwise aggregation",rect_for(panels[0]),[
      [["ss","ν","t","max"]," = max₍₁≤d≤p₎ ",["ss","ν","t","(d)"]],
      [["ss","ν","t","avg"]," = ",["frac",["1"],["p"]]," ∑ᵈ₌₁ᵖ ",["ss","ν","t","(d)"]],
    ],24)
    add_math(s8,"Code-grounded LOESS and peak estimates",rect_for(panels[1]),[
      ["ν̃ₜᵐᵃˣ = LOESS₀.₁(νₜᵐᵃˣ)"],
      ["ν̃ₜᵃᵛᵍ = LOESS₀.₁(νₜᵃᵛᵍ)"],
      ["bₘₐₓ = argmaxₜ ν̃ₜᵐᵃˣ"],
      ["bₐᵥg = argmaxₜ ν̃ₜᵃᵛᵍ"],
    ],19)
    assert all(text(s)==old for s,old in zip(deck.slides,original_text))
    for i,slide in enumerate(deck.slides):
        if i not in (6,7):assert media(slide)==other_media[i],i
    assert media(s8)==other_media[7]
    assert len(other_media[6])==len(media(s7))+1
    assert len(s7._element.xpath('.//*[local-name()="oMath"]'))==4
    assert len(s8._element.xpath('.//*[local-name()="oMath"]'))==6
    dst.parent.mkdir(parents=True,exist_ok=True)
    # Modify exactly TWO slide XML members in the original OOXML ZIP package.
    # This avoids python-pptx rewriting unrelated package parts or dropping
    # unsupported objects, rich metadata, timings, assets, and relationships.
    replacements={
       "ppt/slides/slide7.xml":etree.tostring(s7._element,encoding="UTF-8",xml_declaration=True,standalone=True),
       "ppt/slides/slide8.xml":etree.tostring(s8._element,encoding="UTF-8",xml_declaration=True,standalone=True),
    }
    with zipfile.ZipFile(src,"r") as zin,zipfile.ZipFile(dst,"w") as zout:
        for entry in zin.infolist():
            zout.writestr(entry,replacements.get(entry.filename,zin.read(entry.filename)))
    with zipfile.ZipFile(src) as zin,zipfile.ZipFile(dst) as zout:
        before={name:hashlib.sha256(zin.read(name)).hexdigest() for name in zin.namelist() if name not in replacements}
        after={name:hashlib.sha256(zout.read(name)).hexdigest() for name in zout.namelist() if name not in replacements}
        assert before==after,"An unrelated PowerPoint package part was modified"
        print("VERIFIED every package part outside slides 7 and 8 retains original SHA256")
    check=Presentation(str(dst))
    assert len(check.slides)==44
    assert [len(check.slides[i]._element.xpath('.//*[local-name()="oMath"]')) for i in (6,7)]==[4,6]
    assert all(text(slide)==old for slide,old in zip(check.slides,original_text))
    print("VERIFIED: native editable OMML equations 4 on slide 7, 6 on slide 8")
    print("VERIFIED: no new equation image, 44 slides, original text and other visual media unchanged")
if __name__=="__main__":
    repair(Path(sys.argv[1]),Path(sys.argv[2]))
