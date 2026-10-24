"""Audit an existing PowerPoint deck without changing it.

Outputs a slide-by-slide text inventory and candidate overlap pairs. Numerical
similarity identifies REVIEW candidates only; never delete slides automatically.
Usage: python scripts/audit_initial_assessment_slides.py
"""
from collections import Counter, defaultdict
from difflib import SequenceMatcher
from pathlib import Path
import re

from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "presentations/EWMA_Covariate_Shift_Progress_Review_Haocheng_Lin.pptx"
REPORT = ROOT / "presentations/initial_assessment_slide_inventory.md"

def texts(slide):
    parts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            s = shape.text.strip()
            if s:
                parts.append(s)
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    t = cell.text.strip()
                    if t:
                        parts.append(t)
    return parts

def norm(text):
    return re.sub(r"\\s+", " ", re.sub(r"[^a-z0-9 ]+", " ", text.lower())).strip()

def main():
    prs = Presentation(SOURCE)
    info = [texts(slide) for slide in prs.slides]
    out = [
        "# Initial assessment presentation: slide inventory",
        "",
        f"Source: \`{SOURCE.relative_to(ROOT)}\` (preserved without changes).",
        f"Slides: **{len(info)}**.",
        "",
        "This is a text-layer audit. Text embedded in images or editable maths objects",
        "may not appear. Similarity is a triage signal, NOT proof of duplication.",
        "",
        "## Slide-by-slide content",
        "",
    ]
    for i, group in enumerate(info, 1):
        title = group[0].replace("\\n", " ")[:180] if group else "(No ordinary text title)"
        out += [f"### Slide {i}: {title}", "", f"- Text shapes/tables: {len(group)}"]
        for s in group[:12]:
            line = s.replace("\\n", " / ").replace("|", "\\|")
            out += [f"- {line[:350]}"]
        if len(group)>12: out += [f"- ... plus {len(group)-12} additional text objects"]
        out += [""]
    pairs = []
    for i, a in enumerate(info):
        aa = norm(" ".join(a))
        if len(aa)<70: continue
        for j, b in enumerate(info[:i]):
            bb = norm(" ".join(b))
            if len(bb)<70: continue
            ratio = SequenceMatcher(None, aa[:2500], bb[:2500]).ratio()
            if ratio>=0.54:
                pairs.append((ratio,j+1,i+1))
    out += ["## Potentially overlapping slide pairs (review manually)", ""]
    for similarity,a,b in sorted(pairs, reverse=True)[:45]:
        out += [f"- Slides **{a} / {b}**: text similarity {similarity:.0%}"]
    if not pairs:out += ["- No pairs above the text similarity threshold."]
    out += ["", "## Recurrent opening text", ""]
    occurrences=defaultdict(set)
    for i,g in enumerate(info,1):
        for s in g:
            for line in s.splitlines():
                t=norm(line)
                if 25<=len(t)<=250:occurrences[t].add(i)
    repeats=sorted(((len(ids),s,sorted(ids)) for s,ids in occurrences.items() if len(ids)>=2),reverse=True)
    for count,s,ids in repeats[:40]:
        out += [f"- {count} slides {ids}: {s[:160]}"]
    if not repeats:out+=["- No recurrent text lines."]
    REPORT.write_text("\\n".join(out)+"\\n",encoding="utf-8")
    print(f"Audited {len(info)} slides; potential overlap pairs {len(pairs)}; report at {REPORT}")
    for i,g in enumerate(info,1):
        print(f"SLIDE {i:02d}: {(g[0] if g else '(no text)').replace(chr(10),' / ')[:140]}")
    for ratio,a,b in sorted(pairs, reverse=True)[:18]:
        print(f"OVERLAP {a:02d}/{b:02d}: {ratio:.0%}")

if __name__ == "__main__":
    main()
