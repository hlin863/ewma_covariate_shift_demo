#!/usr/bin/env python3
"""Restore slide 8's Park aggregation/smoothing equations from cited paper.

The earlier equation-repair routine preserves original embedded image bytes.
These two slide-8 panels have NO embedded equations in the October deck, so
the formula illustrations are explicitly *reconstructed* from Park et al.
(2023) Section 2.1 / Algorithm 1, and the repository's park_mean_cusum.py.
Never overwrite the source research PPTX.
"""
from __future__ import annotations
import io
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches

PANELS = (
    (2, 3, (
        r"$\nu_t^{\max} = \max_{1\leq d\leq p}\nu_t^{(d)}$",
        r"$\nu_t^{\mathrm{avg}} = \frac{1}{p}\sum_{d=1}^{p}\nu_t^{(d)}$",
    )),
    (4, 5, (
        r"$\tilde{\nu}_t^{\max}=\mathcal{S}_{\alpha}(\nu_t^{\max})$",
        r"$\tilde{\nu}_t^{\mathrm{avg}}=\mathcal{S}_{\alpha}(\nu_t^{\mathrm{avg}})$",
        r"$b_{\max}=\mathrm{arg\,max}_t\tilde{\nu}_t^{\max}$",
        r"$b_{\mathrm{avg}}=\mathrm{arg\,max}_t\tilde{\nu}_t^{\mathrm{avg}}$",
    )),
)

def render_formula(lines: tuple[str, ...]) -> bytes:
    fig=plt.figure(figsize=(11,3.3),dpi=180,facecolor="#ffffff")
    fig.patch.set_alpha(0.96)
    for j,line in enumerate(lines):
        fig.text(.5, 1.0-(j+1)/(len(lines)+1), line,
                 ha="center",va="center",fontsize=25 if len(lines)==2 else 21,
                 color="#152942")
    bio=io.BytesIO()
    fig.savefig(bio,format="png",bbox_inches="tight",pad_inches=.20,dpi=180,
                facecolor=fig.get_facecolor())
    plt.close(fig)
    return bio.getvalue()

def add_formulas(source:Path, output:Path)->None:
    if source.resolve()==output.resolve():
        raise ValueError("Refusing to overwrite the PPTX source.")
    deck=Presentation(str(source))
    if len(deck.slides)<8:
        raise ValueError("Not enough slides")
    slide=deck.slides[7]
    all_text=lambda s:"\n".join(sh.text for sh in s.shapes if sh.has_text_frame)
    before=all_text(slide)
    title=before.lower()
    if "channelwise mean-cusum" not in title or "smooth" not in title:
        raise ValueError("Slide 8 does not match the expected Park method slide.")
    if any(sh.shape_type==13 for sh in slide.shapes):
        raise ValueError("Slide 8 already contains pictures; avoid adding duplicate equations.")
    for group_idx,label_idx,equations in PANELS:
        group=slide.shapes[group_idx]
        label=slide.shapes[label_idx]
        left=group.left+Inches(.30)
        width=group.width-Inches(.60)
        top=max(group.top+Inches(.40), label.top+label.height+Inches(.23))
        available_bottom=group.top+group.height-Inches(.35)
        height=min(Inches(2.65),available_bottom-top)
        if height<Inches(1.75):
            raise ValueError(f"Insufficient space for slide 8 formulas: {height/914400:.2f} in")
        top=top+(available_bottom-top-height)//2
        slide.shapes.add_picture(io.BytesIO(render_formula(equations)),left,top,width,height)
    assert all_text(slide)==before,"Original explanatory text was altered."
    output.parent.mkdir(parents=True,exist_ok=True)
    deck.save(str(output))
    proof=Presentation(str(output)).slides[7]
    assert all_text(proof)==before
    pictures=[sh for sh in proof.shapes if sh.shape_type==13]
    if len(pictures)!=2:
        raise AssertionError("Both mathematical illustration images must survive round-trip")
    print("VERIFIED: slide 8 now contains TWO source-grounded formula panels without rewriting its text")
    for shape in pictures:
        print("FORMULA_IMAGE",len(shape.image.blob),
              [round(q/914400,2) for q in (shape.left,shape.top,shape.width,shape.height)])

if __name__=="__main__":
    import argparse
    p=argparse.ArgumentParser()
    p.add_argument("source",type=Path)
    p.add_argument("output",type=Path)
    args=p.parse_args()
    add_formulas(args.source,args.output)
