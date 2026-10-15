"""Read-only forensic inspection of slides 7 and 8 across Git revisions."""
import hashlib,io,subprocess,sys
from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
PPTX="presentations/EWMA_Covariate_Shift_Progress_Review_Haocheng_Lin.pptx"
OUT=Path("presentations/formula_integrity_diagnostics")
OUT.mkdir(parents=True,exist_ok=True)
for ref in sys.argv[1:]:
 print("===== REVISION",ref,"=====")
 try: blob=subprocess.check_output(["git","show",f"{ref}:{PPTX}"])
 except subprocess.CalledProcessError: print("UNAVAILABLE",ref);continue
 prs=Presentation(io.BytesIO(blob))
 print("BYTES",len(blob),"SHA256",hashlib.sha256(blob).hexdigest(),"SLIDES",len(prs.slides))
 for idx in (6,7):
  if idx>=len(prs.slides):continue
  slide=prs.slides[idx]
  texts=[sh.text for sh in slide.shapes if sh.has_text_frame]
  omml=len(slide._element.xpath('.//*[local-name()="oMath"]'))
  print("SLIDE",idx+1,"TITLE",repr(texts[0][:130] if texts else ""),"OMML",omml,"SHAPES",len(slide.shapes))
  for k,sh in enumerate(slide.shapes):
   print("  ITEM",k,sh.name,str(sh.shape_type),"TEXT",repr(sh.text[:400] if sh.has_text_frame else ""))
   group=list(sh.shapes) if sh.shape_type==MSO_SHAPE_TYPE.GROUP else [sh]
   for j,ch in enumerate(group):
    for q,blip in enumerate(ch._element.xpath('.//*[local-name()="blip"]')):
     rid=blip.get(qn("r:embed"))
     if not rid:continue
     try: img=slide.part.related_part(rid).blob
     except Exception as e: print("  RELATION ERROR",str(e));continue
     ext="jpg" if img[:3]==b"\xff\xd8\xff" else "png" if img[:4]==b"\x89PNG" else "bin"
     dest=OUT/f"{ref.replace('/','_')}_slide{idx+1}_shape{k}_{j}_{q}.{ext}"
     dest.write_bytes(img)
     print("    IMAGE",dest.name,"BYTES",len(img),"SHA256",hashlib.sha256(img).hexdigest())
