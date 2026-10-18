"""Rebuild only physical slide 8 in the original 44-slide research presentation.

Usage: python -m scripts.update_slide8_park --input SOURCE.pptx --output OUT.pptx
All seven equations are native editable OfficeMath, not pictures.
All explanatory labels remain editable PowerPoint text.
"""
from __future__ import annotations
import argparse
from pathlib import Path
from pptx import Presentation
from scripts.repair_slide8_native_math import add_equation, count_math
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

BG = RGBColor(247,249,252)
NAVY = RGBColor(21,38,62)
TEAL = RGBColor(0,148,138)
SECONDARY = RGBColor(70,89,114)
PALE_BLUE = RGBColor(213,226,238)

def text(slide, message, x, y, w, h, *, size=15, color=NAVY, bold=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    p = frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    p.space_after = Pt(0)
    run = p.add_run()
    run.text = message
    run.font.name = 'Aptos'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape

def equation(slide, latex, x, y, w, h, name):
    """Add source-mapped native OfficeMath, never an equation image."""
    return add_equation(slide, name, Inches(x), Inches(y), Inches(w), Inches(h),
                        font_pt=18)

def panel(slide,x,y,w,h,heading,detail,eqs,label):
    p=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                             Inches(x),Inches(y),Inches(w),Inches(h))
    p.name='Slide 8 / '+label+' / panel'
    p.fill.solid()
    p.fill.fore_color.rgb=RGBColor(255,255,255)
    p.line.color.rgb=PALE_BLUE
    p.line.width=Pt(1.2)
    try:
        p.adjustments[0]=0.15
    except (IndexError,ValueError):
        pass
    text(slide,heading,x+.25,y+.18,w-.48,.34,size=16.2,color=TEAL,bold=True).name='Slide 8 / '+label+' / heading'
    for i,eq in enumerate(eqs):
        equation(slide,eq,x+.30,y+.62+(i*.44)+(.06 if len(eqs)==1 else 0),
                 w-.60,.40,'Slide 8 / '+label+' / equation '+str(i+1))
    text(slide,detail,x+.27,y+h-.49,w-.55,.36,
         size=11.7,color=SECONDARY).name='Slide 8 / '+label+' / explanation'

def edit(source:Path,target:Path):
    prs=Presentation(str(source))
    if len(prs.slides)!=44:
        raise ValueError('Expected exact 44-slide main presentation, found '+str(len(prs.slides)))
    # The actual 44-slide repository deck uses a custom page size.
    # Author the design on a 13.333 x 7.5 canvas, then scale to match it.
    if prs.slide_width <= 0 or prs.slide_height <= 0:
        raise ValueError('Invalid presentation dimensions')
    original=[[sh.text for sh in sl.shapes if sh.has_text_frame] for sl in prs.slides]
    sl=prs.slides[7]
    for shape in list(sl.shapes):
        sl.shapes._spTree.remove(shape._element)
    sl.background.fill.solid()
    sl.background.fill.fore_color.rgb=BG
    text(sl,'From Mean-CUSUM to Mean-Change Confirmation',.82,.66,11.9,.55,
         size=29.8,bold=True).name='Slide 8 / title'
    text(sl,'PARK ET AL. (2023)  /  §2.3, ALGORITHM 1, STEPS 1–4',
         .83,1.28,11.8,.25,size=12.1,color=TEAL,bold=True).name='Slide 8 / subtitle'
    panel(sl,.78,1.70,5.82,2.12,
          '01   Aggregation Across Channels',
          'Maximum and average combine channelwise mean-CUSUM evidence.',
          [r'\nu_t^{\max}=\max_{1\leq d\leq p}\nu_t^{(d)}',
           r'\nu_t^{\mathrm{avg}}=\frac{1}{p}\sum_{d=1}^{p}\nu_t^{(d)}'],'step1')
    panel(sl,6.74,1.70,5.82,2.12,
          '02   Smooth and Locate Peaks',
          'Local polynomial span 0.1; peak locations from smoothed curves.',
          [r'\tilde{\nu}_t^{g}=\mathrm{LP}_{0.1}(\nu^{g})(t)',
           r'b_g=\arg\max_t\tilde{\nu}_t^{g}'],'step2')
    panel(sl,.78,4.00,5.82,2.12,
          '03   Check Unimodality',
          'k = 10 interior positions per side; at least one A count must be 0.',
          [r'A^{\max}A^{\mathrm{avg}}=0'],'step3')
    panel(sl,6.74,4.00,5.82,2.12,
          '04   Confirm and Estimate',
          'If both tests pass, estimate the mean change; else test second order.',
          [r'\left|b^{\max}-b^{\mathrm{avg}}\right|<\frac{\sqrt{T}}{2}',
           r'\hat{b}=\frac{b^{\max}+b^{\mathrm{avg}}}{2}'],'step4')
    text(sl,'Implemented in Python: Steps 1–4  |  Statistical calibration & second-order detection remain open.',
         .94,6.36,11.9,.34,size=13,color=NAVY,bold=True).name='Slide 8 / status'
    text(sl,'Source: Park et al. (2023), §2.3, Algorithm 1  |  Code: src/detection/stage1/park_mean_cusum.py',
         .83,6.79,11.95,.22,size=9.3,color=SECONDARY).name='Slide 8 / citation'
    text(sl,'Research progress  |  hlin863/ewma_covariate_shift_demo  |  Oct 2026',
         .82,7.16,11.8,.18,size=8.9,color=SECONDARY).name='Slide 8 / footer'
    # Scale every new shape to the source deck's *actual* canvas, including
    # text sizes. This keeps the four-card layout proportional to the deck.
    scale_x = prs.slide_width / Inches(13.3333333333)
    scale_y = prs.slide_height / Inches(7.5)
    for shape in sl.shapes:
        shape.left = int(shape.left * scale_x)
        shape.top = int(shape.top * scale_y)
        shape.width = int(shape.width * scale_x)
        shape.height = int(shape.height * scale_y)
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if run.font.size:
                        run.font.size = Pt(run.font.size.pt * scale_y)
        if shape.name.find("EDITABLE OfficeMath equation") >= 0:
            for end in shape._element.xpath('.//*[local-name()="endParaRPr"]'):
                sz=end.get("sz")
                if sz is not None:
                    end.set("sz", str(round(int(sz)*scale_y)))
    print('Source presentation canvas:', round(prs.slide_width/914400,3),
          'x', round(prs.slide_height/914400,3), 'inches')
    if any([s.text for s in prs.slides[i].shapes if s.has_text_frame]!=original[i]
            for i in range(len(prs.slides)) if i!=7):
        raise AssertionError('Non-slide-8 text changed')
    target.parent.mkdir(parents=True,exist_ok=True)
    prs.save(str(target))
    check=Presentation(str(target))
    if len(check.slides)!=44 or len(check.slides[7].shapes)!=24:
        raise AssertionError('Slide-count or shape-count verification failed')
    if count_math(check.slides[7]) != 7 or any(
            sh.shape_type == 13 for sh in check.slides[7].shapes):
        raise AssertionError('Slide 8 must have 7 native OMML equations, 0 pictures')
    print('Slide 8 replaced in 44-slide deck; other 43 slide text sets preserved')
    print('Seven genuine editable OfficeMath expressions and explanatory text verified')
    print('Saved:',target)

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input',type=Path,required=True)
    parser.add_argument('--output',type=Path,required=True)
    args=parser.parse_args()
    if args.input.resolve()==args.output.resolve():
        parser.error('Use different input/output paths to avoid overwriting source')
    edit(args.input,args.output)

if __name__=='__main__':
    main()
